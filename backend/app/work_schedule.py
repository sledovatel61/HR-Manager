"""«График выхода на работу» (Phase 18) — общая логика.

Модуль собирает дневной график выходов из двух источников:

* кандидаты с заполненной датой выхода (``candidates.start_date``);
* служебные строки без кандидата (``schedule_entries``) — «Увольнение
  13:00–14:00», «перевод», «отработка грузчик», «медосмотр».

Правила, зафиксированные здесь (и проверяемые тестами):

* soft-deleted кандидаты не попадают в график никогда;
* ``rejected`` («Отказ») показывается с пометкой «не вышел» и по умолчанию
  отфильтрован — его включает ``include_rejected=true`` (или явный фильтр
  ``stage=rejected``);
* ``fired`` («Уволен») остаётся в графике с пометкой «уволен»: строка не
  исчезает молча;
* порядок строк — дата, затем время (строки без времени идут в конце дня),
  затем ФИО/текст; нумерация внутри дня считается один раз на сервере, чтобы
  экран, Excel и печать совпадали;
* права переиспользуют существующую модель: HR видит своих кандидатов
  (``owner_user_id``), руководитель/администратор — всех; активный грант
  ``candidate_documents_all`` или ``pilot_full_access`` расширяет HR до всей
  базы. Новой модели прав не вводится.
"""

from __future__ import annotations

import io
from collections.abc import Iterable, Sequence
from datetime import date, datetime, time
from typing import Any, cast
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    AccessGrantScope,
    Candidate,
    CandidateSource,
    CandidateStage,
    ScheduleEntry,
    ScheduleImport,
    ScheduleImportRow,
    User,
    UserRole,
)
from app.schemas import (
    ScheduleEntryOut,
    WorkScheduleRow,
)

# Этапы, на которых разрешено назначать дату выхода.
START_STAGES: tuple[CandidateStage, ...] = (
    CandidateStage.OFFER,
    CandidateStage.HIRED,
    CandidateStage.STARTED,
)

# Текстовые поля блока «Выход на работу»: значение записывается только если
# поле реально пришло в запросе (см. ``model_fields_set``), поэтому null
# означает «очистить», а отсутствие поля — «не менять».
START_TEXT_FIELDS: tuple[str, ...] = (
    "start_organization",
    "start_department",
    "shift",
    "start_comment",
)
START_DATE_FIELDS: tuple[str, ...] = ("start_date", "start_time")

# Пометки строк. ``rejected`` — «не вышел», ``fired`` — «уволен»: строка не
# пропадает молча, но и не выглядит как планируемый выход.
_STARTED_STATUS = ("started", "Вышел")
_NOT_CAME_STATUS = ("not_came", "Не вышел")
_DISMISSED_STATUS = ("dismissed", "Уволен")
_PLANNED_STATUS = ("planned", "Планируется")

_STATUS_BY_STAGE: dict[CandidateStage, tuple[str, str]] = {
    CandidateStage.STARTED: _STARTED_STATUS,
    CandidateStage.REJECTED: _NOT_CAME_STATUS,
    CandidateStage.FIRED: _DISMISSED_STATUS,
}

# Пометка внутри колонки «ФИО» — интерфейс и Excel показывают одно и то же.
_NAME_NOTES: dict[str, str] = {
    "not_came": "не вышел",
    "dismissed": "уволен",
}

_EXCEL_HEADERS = (
    "№",
    "ФИО",
    "Время",
    "Организация",
    "Отдел",
    "Должность",
    "Смена",
    "Комментарий",
    "Ответственный HR",
)
_EXCEL_WIDTHS = (5, 34, 12, 20, 26, 24, 11, 34, 20)

_WEEKDAYS = ("пн", "вт", "ср", "чт", "пт", "сб", "вс")
_MONTHS = (
    "янв.",
    "февр.",
    "мар.",
    "апр.",
    "мая",
    "июн.",
    "июл.",
    "авг.",
    "сент.",
    "окт.",
    "нояб.",
    "дек.",
)

# CSV/формулы-инъекции: значения, начинающиеся с этих символов, обязаны
# попадать в файл как текст (а не как формула).
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def day_label(day: date) -> str:
    """«пн, 10 авг. 2026» — формат строки-разделителя из образца."""
    return f"{_WEEKDAYS[day.weekday()]}, {day.day} {_MONTHS[day.month - 1]} {day.year}"


def day_label_short(day: date) -> str:
    """«10.08.2026» — компактная дата для подписей и имён файлов."""
    return f"{day.day:02d}.{day.month:02d}.{day.year:04d}"


def hr_sees_all_candidates(db: Session, user: User) -> bool:
    """Может ли пользователь видеть кандидатов вне своей очереди.

    Руководитель и администратор — да. HR — только при активном гранте
    ``candidate_documents_all`` (scope «все кандидаты») или
    ``pilot_full_access``. Проверка переиспользует ``app.documents.has_grant``,
    чтобы не дублировать модель прав.
    """

    # Identity-map копия могла устареть (роль/грант менялись в другой сессии).
    fresh = db.get(User, user.id, populate_existing=True) or user
    if fresh.role != UserRole.HR:
        return True
    from app.documents import has_grant

    return has_grant(db, fresh, AccessGrantScope.CANDIDATE_DOCUMENTS_ALL)


def visible_candidate_conditions(db: Session, user: User) -> list[Any]:
    """Условия видимости кандидатов (та же модель, что у карточки)."""

    if hr_sees_all_candidates(db, user):
        return []
    return [Candidate.owner_user_id == user.id]


def can_manage_entry(db: Session, user: User, entry: ScheduleEntry) -> bool:
    """Может ли пользователь править/удалять конкретную служебную строку.

    Служебная строка — общий элемент графика, но не «ничей»: её меняет автор
    записи, а также те, кто и так видит всех кандидатов (руководитель,
    администратор, HR с активным грантом «все кандидаты» /
    ``pilot_full_access``). Новой модели прав не вводится — переиспользуется
    ``hr_sees_all_candidates``; строка без автора (автор деактивирован,
    ``ON DELETE SET NULL``) остаётся доступной только этим ролям и грантам.
    """

    if entry.author_user_id is not None and entry.author_user_id == user.id:
        return True
    return hr_sees_all_candidates(db, user)


def _cleaned(value: str | None) -> str:
    return (value or "").strip().casefold()


def _matches(value: str | None, needle: str) -> bool:
    return needle in _cleaned(value)


def build_rows(
    db: Session,
    *,
    user: User,
    date_from: date | None = None,
    date_to: date | None = None,
    organization: str | None = None,
    department: str | None = None,
    position: str | None = None,
    shift: str | None = None,
    owner_id: UUID | None = None,
    stage: CandidateStage | None = None,
    query: str | None = None,
    include_rejected: bool = False,
) -> list[WorkScheduleRow]:
    """Собрать, отфильтровать, отсортировать и пронумеровать строки графика.

    Фильтры по подразделениям/должности/смене и поиск выполняются в Python по
    ``casefold``: это единственный способ одинаково фильтровать кириллицу на
    PostgreSQL и на SQLite (SQLite ``lower()`` не сворачивает регистр
    кириллицы). Объём строк ограничен самим смыслом графика — в него попадают
    только кандидаты с назначенной датой выхода.
    """

    scope = visible_candidate_conditions(db, user)
    candidate_stmt = select(Candidate).where(Candidate.deleted_at.is_(None))
    if scope:
        candidate_stmt = candidate_stmt.where(*scope)
    if date_from is not None:
        candidate_stmt = candidate_stmt.where(Candidate.start_date >= date_from)
    if date_to is not None:
        candidate_stmt = candidate_stmt.where(Candidate.start_date <= date_to)
    # HR всегда видит только свои строки; для руководителя/администратора
    # фильтр по ответственному работает явно.
    if owner_id is not None and not scope:
        candidate_stmt = candidate_stmt.where(Candidate.owner_user_id == owner_id)

    candidates_with_dates = list(db.scalars(candidate_stmt).all())
    # Imported candidates whose source row disappeared from the latest master
    # remain in the candidate database, but do not leak back into the current
    # schedule after their source row was marked inactive.
    active_import_candidate_ids = set(
        db.scalars(
            select(ScheduleImportRow.candidate_id).where(
                ScheduleImportRow.is_active.is_(True),
                ScheduleImportRow.row_type == "person",
                ScheduleImportRow.candidate_id.is_not(None),
            )
        ).all()
    )
    # Candidates without a date cannot be placed in a dated day block. They
    # remain visible in the separate current-import rows and in its Excel.
    candidates_with_dates = [
        item
        for item in candidates_with_dates
        if item.start_date is not None
        and (item.source != CandidateSource.EXCEL_IMPORT or item.id in active_import_candidate_ids)
    ]

    entry_stmt = select(ScheduleEntry).where(ScheduleEntry.is_active.is_(True))
    if date_from is not None:
        entry_stmt = entry_stmt.where(ScheduleEntry.entry_date >= date_from)
    if date_to is not None:
        entry_stmt = entry_stmt.where(ScheduleEntry.entry_date <= date_to)
    entries = list(db.scalars(entry_stmt).all())

    rows: list[WorkScheduleRow] = []
    for candidate in candidates_with_dates:
        status, status_label = _STATUS_BY_STAGE.get(candidate.stage, _PLANNED_STATUS)
        rows.append(
            WorkScheduleRow(
                kind="candidate",
                id=candidate.id,
                candidate_id=candidate.id,
                entry_date=candidate.start_date,  # type: ignore[arg-type]
                number=0,
                start_time=candidate.start_time,
                full_name=candidate.full_name,
                display_name=_display_name(candidate.full_name, status),
                organization=candidate.start_organization,
                department=candidate.start_department,
                position=candidate.position,
                shift=candidate.shift,
                comment=candidate.start_comment,
                owner_user_id=candidate.owner_user_id,
                owner_username=candidate.owner_username,
                stage=candidate.stage,
                status=status,  # type: ignore[arg-type]
                status_label=status_label,
            )
        )

    author_ids = {entry.author_user_id for entry in entries if entry.author_user_id is not None}
    authors: dict[UUID, str] = {}
    if author_ids:
        authors = {
            user.id: user.username
            for user in db.scalars(select(User).where(User.id.in_(author_ids))).all()
        }
    for entry in entries:
        rows.append(
            WorkScheduleRow(
                kind="entry",
                id=entry.id,
                candidate_id=None,
                entry_date=entry.entry_date,
                number=0,
                start_time=entry.time_from,
                end_time=entry.time_to,
                full_name=None,
                display_name=entry.title,
                organization=entry.organization,
                department=entry.department,
                position="",
                shift=None,
                comment=entry.comment,
                owner_user_id=entry.author_user_id,
                owner_username=(
                    authors.get(entry.author_user_id) if entry.author_user_id is not None else None
                ),
                stage=None,
                status="planned",
                status_label="",
            )
        )

    needle = _cleaned(query) if query else ""
    selected: list[WorkScheduleRow] = []
    for row in rows:
        if row.kind == "candidate":
            if stage is not None and row.stage != stage:
                continue
            hide_rejected = (
                row.status == "not_came"
                and not include_rejected
                and stage != CandidateStage.REJECTED
            )
            if hide_rejected:
                # Отказавшиеся по умолчанию отфильтрованы (переключатель в UI).
                continue
            if shift and not _matches(row.shift, _cleaned(shift)):
                continue
        else:
            # Служебная строка не относится к этапам воронки и должности:
            # при явном фильтре по ним она не показывается.
            if stage is not None or position:
                continue
            if owner_id is not None and row.owner_user_id != owner_id and not scope:
                continue
        if organization and not _matches(row.organization, _cleaned(organization)):
            continue
        if department and not _matches(row.department, _cleaned(department)):
            continue
        if position and not _matches(row.position, _cleaned(position)):
            continue
        if needle:
            haystack = " ".join(
                part
                for part in (
                    row.full_name,
                    row.display_name,
                    row.organization,
                    row.department,
                    row.position,
                    row.shift,
                    row.comment,
                )
                if part
            )
            if needle not in haystack.casefold():
                continue
        selected.append(row)

    selected.sort(key=_sort_key)
    _number_rows(selected)
    return selected


def _display_name(full_name: str, status: str) -> str:
    note = _NAME_NOTES.get(status)
    return f"{full_name} — {note}" if note else full_name


def _sort_key(row: WorkScheduleRow) -> tuple:
    """Дата → время (без времени — в конец дня) → тип → ФИО/текст → id."""

    has_time = row.start_time is not None
    return (
        row.entry_date,
        not has_time,
        row.start_time or time.min,
        0 if row.kind == "candidate" else 1,
        (row.full_name or row.display_name).casefold(),
        str(row.id),
    )


def _number_rows(rows: list[WorkScheduleRow]) -> None:
    """Сквозная нумерация внутри дня — как в образце «График приемов»."""

    current_day: date | None = None
    number = 0
    for row in rows:
        if row.entry_date != current_day:
            current_day = row.entry_date
            number = 0
        number += 1
        row.number = number


def day_summaries(rows: Sequence[WorkScheduleRow]) -> list[tuple[date, int, int]]:
    """Дни в порядке следования: (дата, число выходов, всего строк)."""

    summary: list[tuple[date, int, int]] = []
    for row in rows:
        if summary and summary[-1][0] == row.entry_date:
            day, planned, total = summary[-1]
            summary[-1] = (
                day,
                planned + (1 if row.kind == "candidate" else 0),
                total + 1,
            )
        else:
            summary.append((row.entry_date, 1 if row.kind == "candidate" else 0, 1))
    return summary


def entry_out(entry: ScheduleEntry) -> ScheduleEntryOut:
    """Ответ CRUD служебной строки (с автором, без утечки лишних полей)."""

    return ScheduleEntryOut(
        id=entry.id,
        entry_date=entry.entry_date,
        time_from=entry.time_from,
        time_to=entry.time_to,
        title=entry.title,
        organization=entry.organization,
        department=entry.department,
        comment=entry.comment,
        author_user_id=entry.author_user_id,
        author_username=entry.author_username or None,
        created_at=entry.created_at,
        updated_at=entry.updated_at,
    )


def entry_audit_details(entry: ScheduleEntry, *, changed: Iterable[str] | None = None) -> str:
    """Строка аудита служебной записи (без персональных данных)."""

    if changed is not None:
        return "; ".join(changed) if changed else "без изменений"
    return f"date={entry.entry_date.isoformat()} time={format_time_range(entry)}"


def format_time_range(entry: ScheduleEntry) -> str:
    """«13:00–14:00», «13:00» или «—» для служебной строки."""

    if entry.time_from is None:
        return "—"
    start = entry.time_from.strftime("%H:%M")
    if entry.time_to is None:
        return start
    return f"{start}–{entry.time_to.strftime('%H:%M')}"


# --- Excel ------------------------------------------------------------------


def _decode_source_cell(value: object) -> object:
    """Restore the small tagged JSON values kept from the imported sheet."""

    if not isinstance(value, dict) or "__type__" not in value:
        return value
    raw = str(value.get("value", ""))
    kind = value.get("__type__")
    try:
        if kind == "datetime":
            return datetime.fromisoformat(raw)
        if kind == "date":
            return date.fromisoformat(raw)
        if kind == "time":
            return time.fromisoformat(raw)
    except ValueError:
        return raw
    return raw


def build_import_source_xlsx(
    db: Session,
    import_record: ScheduleImport,
    rows: Sequence[ScheduleImportRow],
) -> bytes:
    """Build a normalized copy of the latest source table, not a DB dump.

    Source columns and their order are retained, current work fields are
    refreshed from the linked candidate, and an HR-responsible column is
    populated (or appended). Every active person row is exported even when
    there is no date, candidate match, or owner yet.
    """

    workbook = Workbook()
    sheet = cast(Worksheet, workbook.active)
    title = (import_record.sheet_title or "График выхода")[:31] or "График выхода"
    sheet.title = title

    headers = list(import_record.source_headers or [])
    columns = dict(import_record.source_columns or {})
    if not headers:
        # Backward-compatible normalized schema for provenance created before
        # source columns were retained (pre-0021 installations).
        headers = [
            "№",
            "ФИО",
            "Дата и время",
            "Организация",
            "Отдел",
            "Должность",
            "Смена / комментарий",
        ]
        columns = {
            "name": 2,
            "time": 3,
            "organization": 4,
            "department": 5,
            "position": 6,
            "comment": 7,
            "owner": 8,
        }
    max_column = max(
        [len(headers), *(value for value in columns.values() if isinstance(value, int))],
        default=0,
    )
    max_column = max(1, max_column)
    if len(headers) < max_column:
        headers.extend([None] * (max_column - len(headers)))
    owner_column = columns.get("owner")
    if not isinstance(owner_column, int) or owner_column < 1:
        owner_column = max_column + 1
        headers.append("Ответственный HR")
        max_column += 1
    elif owner_column > len(headers):
        headers.extend([None] * (owner_column - len(headers)))
        headers[owner_column - 1] = "Ответственный HR"
        max_column = max(max_column, owner_column)
    else:
        # Keep a user-provided heading if one exists; otherwise name the added
        # responsibility column clearly in the source position.
        if not headers[owner_column - 1]:
            headers[owner_column - 1] = "Ответственный HR"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for index in range(1, max_column + 1):
        cell = sheet.cell(row=1, column=index)
        header_value = headers[index - 1] if index <= len(headers) else None
        _write_text(cell, header_value or "")
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = (
            18 if index > 9 else _EXCEL_WIDTHS[index - 1]
        )
    sheet.row_dimensions[1].height = 30

    name_column = columns.get("name")
    time_column = columns.get("time")
    organization_column = columns.get("organization")
    department_column = columns.get("department")
    position_column = columns.get("position")
    comment_column = columns.get("comment")

    candidate_ids = {row.candidate_id for row in rows if row.candidate_id is not None}
    candidates = (
        {
            candidate.id: candidate
            for candidate in db.scalars(
                select(Candidate).where(Candidate.id.in_(candidate_ids))
            ).all()
        }
        if candidate_ids
        else {}
    )
    entry_ids = {row.entry_id for row in rows if row.entry_id is not None}
    entries = (
        {
            entry.id: entry
            for entry in db.scalars(
                select(ScheduleEntry).where(ScheduleEntry.id.in_(entry_ids))
            ).all()
        }
        if entry_ids
        else {}
    )
    owner_ids = {row.owner_user_id for row in rows if row.owner_user_id is not None}
    owner_ids.update(
        candidate.owner_user_id
        for candidate in candidates.values()
        if candidate.owner_user_id is not None
    )
    owners = (
        {user.id: user for user in db.scalars(select(User).where(User.id.in_(owner_ids))).all()}
        if owner_ids
        else {}
    )

    output_row = 2
    for source in sorted(rows, key=lambda item: (item.row_order, item.sheet_row, str(item.id))):
        if source.row_type not in ("person", "service"):
            continue
        values: list[object] = [
            _decode_source_cell(value) for value in (source.source_values or [])
        ]
        if len(values) < max_column:
            values.extend([None] * (max_column - len(values)))
        elif len(values) > max_column:
            values = values[:max_column]
        candidate = candidates.get(source.candidate_id) if source.candidate_id else None
        linked_entry = entries.get(source.entry_id) if source.entry_id else None
        current_entry = (
            linked_entry
            if source.row_type == "service" and linked_entry is not None and linked_entry.is_active
            else None
        )

        def set_value(
            target_values: list[object],
            column: object,
            value: object,
            *,
            overwrite_empty: bool = True,
        ) -> None:
            if not isinstance(column, int) or column < 1 or column > max_column:
                return
            if not overwrite_empty and value is None:
                return
            target_values[column - 1] = value

        # Keep the source spelling for people. Linked service entries may be
        # edited from the schedule page, so export their live title and fields.
        if source.row_type == "service" and current_entry is not None:
            set_value(values, name_column, current_entry.title)
        elif source.full_name:
            set_value(values, name_column, source.full_name)

        effective_date = (
            candidate.start_date
            if candidate is not None
            else current_entry.entry_date
            if current_entry is not None
            else source.entry_date
        )
        effective_time = (
            candidate.start_time
            if candidate is not None
            else current_entry.time_from
            if current_entry is not None
            else source.time_from
        )
        effective_time_to = current_entry.time_to if current_entry is not None else source.time_to
        if effective_date is not None and effective_time_to is not None:
            start_text = effective_time.strftime("%H:%M") if effective_time else "—"
            time_value: object = (
                f"{effective_date.strftime('%d.%m.%Y')} к {start_text}–"
                f"{effective_time_to.strftime('%H:%M')}"
            )
        elif effective_date is not None:
            time_value = datetime.combine(effective_date, effective_time or time.min)
        else:
            time_value = effective_time
        set_value(values, time_column, time_value)

        organization = (
            candidate.start_organization
            if candidate is not None
            else current_entry.organization
            if current_entry is not None
            else source.organization
        )
        department = (
            candidate.start_department
            if candidate is not None
            else current_entry.department
            if current_entry is not None
            else source.department
        )
        position = candidate.position if candidate is not None else (source.position or "")
        set_value(values, organization_column, organization)
        set_value(values, department_column, department)
        set_value(values, position_column, position)

        current_comment = (
            candidate.start_comment
            if candidate is not None
            else current_entry.comment
            if current_entry is not None
            else source.comment
        )
        if current_comment is not None:
            shift = candidate.shift if candidate is not None else source.shift
            comment_value = "; ".join(part for part in (shift, current_comment) if part)
            set_value(values, comment_column, comment_value)
        elif current_entry is not None:
            set_value(values, comment_column, "")
        # An originally populated comment cell may be a phone number. Leave it
        # intact when the parsed comment is empty; raw phone/notes columns are
        # otherwise kept in their original positions.

        owner_id = (
            candidate.owner_user_id
            if candidate is not None and candidate.owner_user_id is not None
            else source.owner_user_id
        )
        owner = owners.get(owner_id) if owner_id is not None else None
        owner_name = (owner.full_name or owner.username) if owner is not None else ""
        set_value(values, owner_column, owner_name)

        for column, cell_value in enumerate(values, start=1):
            cell = sheet.cell(row=output_row, column=column)
            _write_text(cell, cell_value)
            cell.border = border
            if column == time_column and isinstance(cell_value, datetime):
                cell.number_format = (
                    "dd.mm.yyyy hh:mm" if cell_value.time() != time.min else "dd.mm.yyyy"
                )
            if column == time_column and isinstance(cell_value, time):
                cell.number_format = "hh:mm"
        output_row += 1

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(max_column)}{max(1, output_row - 1)}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    if sheet.sheet_properties.pageSetUpPr is not None:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_text(cell: Any, value: Any) -> None:
    """Записать значение так, чтобы «=cmd» не стало формулой."""

    cell.value = value
    if isinstance(value, str) and value[:1] in _FORMULA_PREFIXES:
        # openpyxl по умолчанию помечает строку с ведущим «=» как формулу;
        # принудительный строковый тип делает значение обычным текстом.
        cell.data_type = "s"


def build_xlsx(
    rows: Sequence[WorkScheduleRow],
    *,
    period_from: date | None,
    period_to: date | None,
    filters: dict[str, str],
    include_rejected: bool,
    generated_at_label: str,
) -> bytes:
    """Собрать .xlsx «График выхода на работу» в формате образца.

    Строка-заголовок дня (жирная, с заливкой, «пн, 10 авг. 2026»), под ней
    пронумерованные строки дня; ширина колонок подобрана, печать — альбомная,
    шапка таблицы повторяется на каждой странице. Значения, начинающиеся с
    ``= + - @``, записываются текстом.
    """

    workbook = Workbook()
    sheet = cast(Worksheet, workbook.active)
    sheet.title = "График выхода"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    day_font = Font(bold=True)
    day_fill = PatternFill("solid", fgColor="DCE6F1")
    title_font = Font(bold=True, size=14)
    service_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    last_column = len(_EXCEL_HEADERS)
    last_letter = get_column_letter(last_column)

    _write_text(sheet.cell(row=1, column=1), "График выхода на работу")
    sheet.cell(row=1, column=1).font = title_font
    sheet.merge_cells(f"A1:{last_letter}1")

    if period_from is not None and period_to is not None:
        period = f"Период: {day_label_short(period_from)} – {day_label_short(period_to)}"
    elif period_from is not None:
        period = f"Период: с {day_label_short(period_from)}"
    elif period_to is not None:
        period = f"Период: по {day_label_short(period_to)}"
    else:
        period = "Период: весь график"
    filters_line = ", ".join(f"{name}: {value}" for name, value in filters.items() if value)
    if include_rejected:
        filters_line = "; ".join(part for part in (filters_line, "показаны отказавшиеся") if part)
    subtitle = f"{period}. {filters_line}" if filters_line else period
    _write_text(sheet.cell(row=2, column=1), subtitle)
    sheet.merge_cells(f"A2:{last_letter}2")

    _write_text(sheet.cell(row=3, column=1), f"Сформировано: {generated_at_label}")
    sheet.merge_cells(f"A3:{last_letter}3")

    header_row = 4
    for index, title in enumerate(_EXCEL_HEADERS, start=1):
        cell = sheet.cell(row=header_row, column=index, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for index, width in enumerate(_EXCEL_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    row_index = header_row + 1
    current_day: date | None = None
    day_candidates = 0
    if rows:
        for row in rows:
            if row.entry_date != current_day:
                if current_day is not None:
                    row_index += 1  # пустая строка между днями
                current_day = row.entry_date
                day_candidates = sum(
                    1
                    for item in rows
                    if item.entry_date == current_day and item.kind == "candidate"
                )
                day_cell = sheet.cell(row=row_index, column=1)
                # Строка дня формируется кодом, но пишется тем же санитайзером:
                # defence-in-depth против будущих правок формата/custom-полей.
                _write_text(day_cell, f"{day_label(current_day)} — выходов: {day_candidates}")
                day_cell.font = day_font
                day_cell.fill = day_fill
                day_cell.alignment = Alignment(horizontal="left")
                sheet.merge_cells(
                    start_row=row_index, start_column=1, end_row=row_index, end_column=last_column
                )
                for column in range(1, last_column + 1):
                    sheet.cell(row=row_index, column=column).border = border
                row_index += 1

            time_text = "—"
            if row.start_time is not None:
                time_text = row.start_time.strftime("%H:%M")
            elif row.kind == "entry":
                time_text = "—"
            values: tuple[Any, ...] = (
                row.number,
                row.display_name,
                time_text,
                row.organization or "—",
                row.department or "—",
                row.position or "—",
                row.shift or "—",
                row.comment or "—",
                row.owner_username or "—",
            )
            for column, value in enumerate(values, start=1):
                data_cell = sheet.cell(row=row_index, column=column)
                _write_text(data_cell, value)
                data_cell.border = border
                if row.kind == "entry":
                    data_cell.fill = service_fill
                if column in (1, 3):
                    data_cell.alignment = Alignment(horizontal="center")
            row_index += 1

    if not rows:
        empty = sheet.cell(row=row_index, column=1)
        _write_text(empty, "На выбранный период выходов нет.")
        empty.font = Font(italic=True)
        sheet.merge_cells(
            start_row=row_index, start_column=1, end_row=row_index, end_column=last_column
        )
        row_index += 1
    else:
        row_index += 1
        total_candidates = sum(1 for item in rows if item.kind == "candidate")
        totals = sheet.cell(row=row_index, column=1)
        _write_text(totals, f"Итого выходов за период: {total_candidates}")
        totals.font = day_font

    # Печать: альбомная ориентация, шапка таблицы повторяется на каждой
    # странице, таблица ужимается по ширине страницы.
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    if sheet.sheet_properties.pageSetUpPr is not None:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.print_options.horizontalCentered = False
    sheet.freeze_panes = f"A{header_row + 1}"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
