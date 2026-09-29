"""Unit tests for «График выхода на работу» (phase 18, in-memory SQLite).

Покрытие соответствует промпту этапа:

* карточка: сохранение и перенос даты/времени, очистка полей, 422 при
  «Вышел» без даты, 422 при дате не на финальном этапе, аудит «было → стало»;
* вкладка: права (HR — только свои, руководитель/администратор — все, HR с
  scope «все кандидаты» — все), фильтры, поиск, сортировка по дате и времени
  (без времени — в конце дня), нумерация внутри дня, служебные строки,
  исключение удалённых, `rejected` с пометкой «не вышел» и переключателем;
* CRUD служебных строк и его аудит;
* Excel: файл открывается ``openpyxl.load_workbook``, содержит те же строки,
  что отдаёт API, шапку дня, альбомную печать и повтор шапки, а значения,
  начинающиеся с «=», остаются текстом.

Integration mirror against PostgreSQL: ``tests/test_integration_work_schedule.py``.
"""

from __future__ import annotations

import io
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date, time, timedelta
from typing import cast
from uuid import UUID

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    AccessGrant,
    AccessGrantScope,
    AuditAction,
    AuditEvent,
    Candidate,
    CandidateStage,
    User,
    UserRole,
)
from app.routers.auth import reset_login_limiter
from tests.conftest import FIXTURE_PASSWORD, make_candidate, make_user

MONDAY = date(2026, 8, 10)  # понедельник
TUESDAY = date(2026, 8, 11)


@pytest.fixture(autouse=True)
def _clean_limiter() -> Iterator[None]:
    reset_login_limiter()
    yield
    reset_login_limiter()


def _login(client: TestClient, username: str, password: str = FIXTURE_PASSWORD) -> httpx.Response:
    return client.post("/auth/login", json={"username": username, "password": password})


def _csrf(response: httpx.Response) -> str:
    return response.json()["csrf_token"]


def _clock(value: str) -> time:
    """«09:00» → datetime.time(9, 0) для фикстур."""

    hours, minutes = value.split(":")[:2]
    return time(int(hours), int(minutes))


def _schedule(
    db: Session,
    candidate: Candidate,
    *,
    start_date: date = MONDAY,
    start_time: str | None = None,
    organization: str | None = None,
    department: str | None = None,
    shift: str | None = None,
    comment: str | None = None,
) -> Candidate:
    """Проставить кандидату поля графика напрямую в БД (fixture-хелпер)."""

    candidate.start_date = start_date
    candidate.start_time = _clock(start_time) if start_time else None
    candidate.start_organization = organization
    candidate.start_department = department
    candidate.shift = shift
    candidate.start_comment = comment
    db.commit()
    db.refresh(candidate)
    return candidate


def _iso(value: date) -> str:
    return value.isoformat()


def _sheet(workbook: object) -> Worksheet:
    """Активный лист загруженной книги (stubs openpyxl дают Optional)."""

    return cast(Worksheet, workbook.active)  # type: ignore[attr-defined]


# --- Карточка: дата выхода ---------------------------------------------------


def test_patch_saves_and_moves_start_date_with_audit(
    client: TestClient, db_session: Session
) -> None:
    hr = make_user(db_session, username="hr1", role=UserRole.HR)
    candidate = make_candidate(db_session, owner=hr, stage=CandidateStage.OFFER)
    csrf = _csrf(_login(client, "hr1"))

    saved = client.patch(
        f"/candidates/{candidate.id}",
        json={
            "start_date": _iso(MONDAY),
            "start_time": "09:00",
            "start_organization": "ООО Авион",
            "start_department": "Транзитный склад",
            "shift": "1 смена",
            "start_comment": "при наличии места",
        },
        headers={"X-CSRF-Token": csrf},
    )
    assert saved.status_code == 200
    body = saved.json()
    assert body["start_date"] == _iso(MONDAY)
    assert body["start_time"] == "09:00:00"
    assert body["start_organization"] == "ООО Авион"
    assert body["shift"] == "1 смена"

    # Быстрый перенос даты прямо со вкладки «График выхода»: аудит «было → стало».
    moved = client.patch(
        f"/candidates/{candidate.id}",
        json={"start_date": _iso(TUESDAY), "start_time": "10:30"},
        headers={"X-CSRF-Token": csrf},
    )
    assert moved.status_code == 200
    assert moved.json()["start_date"] == _iso(TUESDAY)

    events = list(
        db_session.scalars(
            select(AuditEvent).where(
                AuditEvent.action == AuditAction.CANDIDATE_START_SCHEDULE_CHANGED
            )
        ).all()
    )
    assert len(events) == 2
    assert "start_date: — -> 2026-08-10" in (events[0].details or "")
    assert "start_date: 2026-08-10 -> 2026-08-11" in (events[1].details or "")
    assert "start_time: 09:00 -> 10:30" in (events[1].details or "")
    assert events[1].actor_user_id == hr.id

    # Очистка: null в явно переданном поле убирает значение из графика.
    cleared = client.patch(
        f"/candidates/{candidate.id}",
        json={"start_date": None, "start_time": None, "start_comment": None},
        headers={"X-CSRF-Token": csrf},
    )
    assert cleared.status_code == 200
    assert cleared.json()["start_date"] is None
    assert cleared.json()["start_time"] is None
    assert cleared.json()["start_comment"] is None
    # Незатронутые поля сохраняются.
    assert cleared.json()["start_organization"] == "ООО Авион"


def test_started_without_date_returns_422_with_russian_text(
    client: TestClient, db_session: Session
) -> None:
    hr = make_user(db_session, username="hr1", role=UserRole.HR)
    candidate = make_candidate(db_session, owner=hr, stage=CandidateStage.OFFER)
    csrf = _csrf(_login(client, "hr1"))

    response = client.patch(
        f"/candidates/{candidate.id}", json={"stage": "started"}, headers={"X-CSRF-Token": csrf}
    )
    assert response.status_code == 422
    assert "дату выхода" in response.json()["detail"].lower()
    db_session.refresh(candidate)
    assert candidate.stage == CandidateStage.OFFER  # ничего не записалось

    # С датой тот же переход проходит.
    ok = client.patch(
        f"/candidates/{candidate.id}",
        json={"stage": "started", "start_date": _iso(MONDAY), "start_time": "08:00"},
        headers={"X-CSRF-Token": csrf},
    )
    assert ok.status_code == 200
    assert ok.json()["stage"] == "started"


def test_start_date_only_on_final_stages(client: TestClient, db_session: Session) -> None:
    hr = make_user(db_session, username="hr1", role=UserRole.HR)
    early = make_candidate(db_session, owner=hr, full_name="Ранний Этап", stage=CandidateStage.NEW)
    offer = make_candidate(db_session, owner=hr, full_name="Оффер", stage=CandidateStage.OFFER)
    csrf = _csrf(_login(client, "hr1"))

    rejected = client.patch(
        f"/candidates/{early.id}",
        json={"start_date": _iso(MONDAY)},
        headers={"X-CSRF-Token": csrf},
    )
    assert rejected.status_code == 422
    assert "Оффер" in rejected.json()["detail"]

    allowed = client.patch(
        f"/candidates/{offer.id}",
        json={"start_date": _iso(MONDAY)},
        headers={"X-CSRF-Token": csrf},
    )
    assert allowed.status_code == 200

    hired = client.patch(
        f"/candidates/{offer.id}", json={"stage": "hired"}, headers={"X-CSRF-Token": csrf}
    )
    assert hired.status_code == 200
    # Перенос на этап «Оформлен» дату сохраняет (её видно в графике).
    assert hired.json()["start_date"] == _iso(MONDAY)

    # Время без даты смысла не имеет — 422 с понятным текстом.
    time_only = client.patch(
        f"/candidates/{early.id}", json={"start_time": "09:00"}, headers={"X-CSRF-Token": csrf}
    )
    assert time_only.status_code == 422
    assert "дату выхода" in time_only.json()["detail"].lower()


def test_hr_cannot_set_start_date_on_foreign_candidate(
    client: TestClient, db_session: Session
) -> None:
    make_user(db_session, username="hr1", role=UserRole.HR)
    hr2 = make_user(db_session, username="hr2", role=UserRole.HR)
    foreign = make_candidate(db_session, owner=hr2, stage=CandidateStage.OFFER)
    csrf = _csrf(_login(client, "hr1"))

    response = client.patch(
        f"/candidates/{foreign.id}",
        json={"start_date": _iso(MONDAY)},
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 404
    db_session.refresh(foreign)
    assert foreign.start_date is None


# --- Вкладка «График выхода»: права, фильтры, сортировка ---------------------


@dataclass(frozen=True)
class _ScheduleData:
    """Общие для тестов данные графика (две HR, кандидаты всех статусов)."""

    hr1: User
    hr2: User
    morning: Candidate
    no_time: Candidate
    evening: Candidate
    refused: Candidate
    dismissed: Candidate
    deleted: Candidate
    fresh: Candidate


def _schedule_fixture(db: Session) -> _ScheduleData:
    hr1 = make_user(db, username="hr1", role=UserRole.HR)
    hr2 = make_user(db, username="hr2", role=UserRole.HR)
    make_user(db, username="mgr", role=UserRole.MANAGER)
    make_user(db, username="adm", role=UserRole.ADMIN)

    morning = make_candidate(db, owner=hr1, full_name="Антонов Антон", stage=CandidateStage.HIRED)
    _schedule(
        db,
        morning,
        start_date=MONDAY,
        start_time="08:00",
        organization="ООО Авион",
        department="Транзитный склад",
        shift="1 смена",
    )
    no_time = make_candidate(db, owner=hr1, full_name="Борисов Борис", stage=CandidateStage.OFFER)
    _schedule(db, no_time, start_date=MONDAY, organization="ИТЦ", department="Отдел кадров")
    evening = make_candidate(db, owner=hr1, full_name="Ветров Виктор", stage=CandidateStage.HIRED)
    _schedule(db, evening, start_date=MONDAY, start_time="18:00", organization="ОП Авион")
    refused = make_candidate(
        db, owner=hr1, full_name="Гончаров Глеб", stage=CandidateStage.REJECTED
    )
    _schedule(db, refused, start_date=MONDAY, start_time="09:00")
    dismissed = make_candidate(db, owner=hr2, full_name="Дроздов Денис", stage=CandidateStage.FIRED)
    _schedule(db, dismissed, start_date=TUESDAY, start_time="07:00", organization="ИТЦ")
    deleted = make_candidate(
        db, owner=hr1, full_name="Ершов Егор", stage=CandidateStage.HIRED, deleted=True
    )
    _schedule(db, deleted, start_date=MONDAY, start_time="11:00")
    fresh = make_candidate(db, owner=hr1, full_name="Жуков Ждан", stage=CandidateStage.NEW)
    return _ScheduleData(
        hr1=hr1,
        hr2=hr2,
        morning=morning,
        no_time=no_time,
        evening=evening,
        refused=refused,
        dismissed=dismissed,
        deleted=deleted,
        fresh=fresh,
    )


def test_schedule_sorts_by_day_and_time_without_time_last(
    client: TestClient, db_session: Session
) -> None:
    data = _schedule_fixture(db_session)
    csrf = _csrf(_login(client, "adm"))
    created = client.post(
        "/work-schedule/entries",
        json={
            "entry_date": _iso(MONDAY),
            "time_from": "13:00",
            "time_to": "14:00",
            "title": "Увольнение",
            "comment": "по заявлению",
        },
        headers={"X-CSRF-Token": csrf},
    )
    assert created.status_code == 201

    response = client.get(f"/work-schedule?from={_iso(MONDAY)}&to={_iso(TUESDAY)}")
    assert response.status_code == 200
    body = response.json()
    names = [item["display_name"] for item in body["items"]]
    # Понедельник по времени (кандидаты и служебные строки вместе), затем без
    # времени; отказавшийся скрыт, удалённый не показывается, уволенный — с
    # пометкой и на своём дне.
    assert names == [
        "Антонов Антон",
        "Увольнение",
        "Ветров Виктор",
        "Борисов Борис",
        "Дроздов Денис — уволен",
    ]
    numbers = [item["number"] for item in body["items"]]
    assert numbers == [1, 2, 3, 4, 1]  # нумерация перезапускается в новом дне
    assert body["days"] == 2
    assert body["total"] == 5
    shown_ids = {item["id"] for item in body["items"]}
    assert str(data.deleted.id) not in shown_ids
    assert str(data.fresh.id) not in shown_ids


def test_rejected_hidden_by_default_and_shown_with_flag(
    client: TestClient, db_session: Session
) -> None:
    data = _schedule_fixture(db_session)
    _login(client, "hr1")

    default = client.get(f"/work-schedule?from={_iso(MONDAY)}&to={_iso(MONDAY)}").json()
    names = [item["full_name"] for item in default["items"]]
    assert "Гончаров Глеб" not in names

    shown = client.get(
        f"/work-schedule?from={_iso(MONDAY)}&to={_iso(MONDAY)}&include_rejected=true"
    ).json()
    refused = next(item for item in shown["items"] if item["full_name"] == "Гончаров Глеб")
    assert refused["status"] == "not_came"
    assert refused["status_label"] == "Не вышел"
    assert refused["display_name"].endswith("— не вышел")

    # Явный фильтр по этапу «Отказ» тоже показывает строку.
    by_stage = client.get(
        f"/work-schedule?from={_iso(MONDAY)}&to={_iso(MONDAY)}&stage=rejected"
    ).json()
    assert [item["full_name"] for item in by_stage["items"]] == ["Гончаров Глеб"]
    assert str(data.refused.id) == by_stage["items"][0]["candidate_id"]


def test_hr_sees_only_own_and_manager_sees_all(client: TestClient, db_session: Session) -> None:
    data = _schedule_fixture(db_session)
    _login(client, "hr1")
    mine = client.get(f"/work-schedule?from={_iso(MONDAY)}&to={_iso(TUESDAY)}").json()
    owners = {item["owner_username"] for item in mine["items"]}
    assert owners == {"hr1"}

    _login(client, "mgr")
    everything = client.get(f"/work-schedule?from={_iso(MONDAY)}&to={_iso(TUESDAY)}").json()
    assert {item["owner_username"] for item in everything["items"]} == {"hr1", "hr2"}

    # Фильтр по ответственному — для руководителя; HR он не расширяет права.
    filtered = client.get(
        f"/work-schedule?from={_iso(MONDAY)}&to={_iso(TUESDAY)}&owner={data.hr1.id}"
    ).json()
    assert {item["owner_username"] for item in filtered["items"]} == {"hr1"}

    _login(client, "hr1")
    scoped = client.get(
        f"/work-schedule?from={_iso(MONDAY)}&to={_iso(TUESDAY)}&owner={data.hr2.id}"
    ).json()
    assert {item["owner_username"] for item in scoped["items"]} == {"hr1"}


def test_hr_with_all_candidates_grant_sees_whole_base(
    client: TestClient, db_session: Session
) -> None:
    data = _schedule_fixture(db_session)
    db_session.add(
        AccessGrant(
            user_id=data.hr1.id,
            scope=AccessGrantScope.CANDIDATE_DOCUMENTS_ALL,
            granted_by_user_id=data.hr1.id,
        )
    )
    db_session.commit()

    _login(client, "hr1")
    body = client.get(f"/work-schedule?from={_iso(MONDAY)}&to={_iso(TUESDAY)}").json()
    assert {item["owner_username"] for item in body["items"]} == {"hr1", "hr2"}


def test_schedule_filters_and_search(client: TestClient, db_session: Session) -> None:
    _schedule_fixture(db_session)
    _login(client, "hr1")
    base = f"/work-schedule?from={_iso(MONDAY)}&to={_iso(TUESDAY)}"

    assert len(client.get(f"{base}&organization=ИТЦ").json()["items"]) == 1
    assert len(client.get(f"{base}&organization=итц").json()["items"]) == 1  # регистр не важен
    assert len(client.get(f"{base}&department=Транзитный склад").json()["items"]) == 1
    assert len(client.get(f"{base}&shift=1 смена").json()["items"]) == 1
    by_position = client.get(f"{base}&position=Слесарь")
    assert by_position.status_code == 200 and by_position.json()["items"] == []

    search = client.get(f"{base}&q=ветров").json()
    assert [item["full_name"] for item in search["items"]] == ["Ветров Виктор"]

    stage = client.get(f"{base}&stage=offer").json()
    assert [item["full_name"] for item in stage["items"]] == ["Борисов Борис"]

    empty = client.get("/work-schedule?from=2026-09-01&to=2026-09-02").json()
    assert empty["items"] == [] and empty["total"] == 0 and empty["days"] == 0

    reversed_period = client.get("/work-schedule?from=2026-09-02&to=2026-09-01")
    assert reversed_period.status_code == 422


def test_suggestions_come_from_entered_values(client: TestClient, db_session: Session) -> None:
    _schedule_fixture(db_session)
    csrf = _csrf(_login(client, "hr1"))
    client.post(
        "/work-schedule/entries",
        json={
            "entry_date": _iso(MONDAY),
            "title": "Медосмотр",
            "organization": "ИТЦ",
            "department": "Поликлиника",
        },
        headers={"X-CSRF-Token": csrf},
    )

    body = client.get("/work-schedule/suggestions").json()
    assert body["organizations"] == ["ИТЦ", "ООО Авион", "ОП Авион"]
    assert "Поликлиника" in body["departments"]
    assert body["shifts"] == ["1 смена"]


def test_hr_cannot_read_foreign_rows_in_detail_lists(
    client: TestClient, db_session: Session
) -> None:
    _schedule_fixture(db_session)
    _login(client, "hr2")
    body = client.get(f"/work-schedule?from={_iso(MONDAY)}&to={_iso(TUESDAY)}").json()
    assert {item["owner_username"] for item in body["items"]} == {"hr2"}


# --- Служебные строки --------------------------------------------------------


def test_schedule_entry_crud_and_audit(client: TestClient, db_session: Session) -> None:
    hr = make_user(db_session, username="hr1", role=UserRole.HR)
    csrf = _csrf(_login(client, "hr1"))

    created = client.post(
        "/work-schedule/entries",
        json={
            "entry_date": _iso(MONDAY),
            "time_from": "13:00",
            "time_to": "14:00",
            "title": "Увольнение",
            "department": "Производственный цех Сокол",
            "comment": "по заявлению",
        },
        headers={"X-CSRF-Token": csrf},
    )
    assert created.status_code == 201
    entry = created.json()
    assert entry["author_username"] == "hr1"
    assert entry["time_from"] == "13:00:00"

    updated = client.patch(
        f"/work-schedule/entries/{entry['id']}",
        json={"time_from": "15:00", "time_to": None, "comment": "перенесено"},
        headers={"X-CSRF-Token": csrf},
    )
    assert updated.status_code == 200
    assert updated.json()["time_from"] == "15:00:00"
    assert updated.json()["time_to"] is None

    broken_interval = client.patch(
        f"/work-schedule/entries/{entry['id']}",
        json={"time_from": "16:00", "time_to": "15:30"},
        headers={"X-CSRF-Token": csrf},
    )
    assert broken_interval.status_code == 422

    missing = client.patch(
        "/work-schedule/entries/00000000-0000-0000-0000-000000000000",
        json={"title": "нет"},
        headers={"X-CSRF-Token": csrf},
    )
    assert missing.status_code == 404

    actions = [
        event.action
        for event in db_session.scalars(
            select(AuditEvent).where(
                AuditEvent.action.in_(
                    [
                        AuditAction.WORK_SCHEDULE_ENTRY_CREATED,
                        AuditAction.WORK_SCHEDULE_ENTRY_UPDATED,
                        AuditAction.WORK_SCHEDULE_ENTRY_DELETED,
                    ]
                )
            )
        ).all()
    ]
    assert AuditAction.WORK_SCHEDULE_ENTRY_CREATED in actions
    assert AuditAction.WORK_SCHEDULE_ENTRY_UPDATED in actions

    deleted = client.delete(f"/work-schedule/entries/{entry['id']}", headers={"X-CSRF-Token": csrf})
    assert deleted.status_code == 204
    assert client.get(f"/work-schedule?from={_iso(MONDAY)}&to={_iso(MONDAY)}").json()["items"] == []
    assert hr.role == UserRole.HR  # автор записи не изменяется


def test_schedule_entry_validation(client: TestClient, db_session: Session) -> None:
    make_user(db_session, username="hr1", role=UserRole.HR)

    # Без сессии ни чтение графика, ни служебные строки недоступны.
    assert client.get("/work-schedule").status_code == 401
    assert (
        client.post(
            "/work-schedule/entries", json={"entry_date": _iso(MONDAY), "title": "Перевод"}
        ).status_code
        == 401
    )

    csrf = _csrf(_login(client, "hr1"))

    blank = client.post(
        "/work-schedule/entries",
        json={"entry_date": _iso(MONDAY), "title": "   "},
        headers={"X-CSRF-Token": csrf},
    )
    assert blank.status_code == 422

    reversed_interval = client.post(
        "/work-schedule/entries",
        json={
            "entry_date": _iso(MONDAY),
            "title": "Перевод",
            "time_from": "14:00",
            "time_to": "13:00",
        },
        headers={"X-CSRF-Token": csrf},
    )
    assert reversed_interval.status_code == 422

    # Мутирующий запрос без CSRF-токена отклоняется (существующая защита).
    assert (
        client.post(
            "/work-schedule/entries",
            json={"entry_date": _iso(MONDAY), "title": "Перевод"},
        ).status_code
        == 403
    )


# --- Excel -------------------------------------------------------------------


def test_export_matches_api_rows_and_is_printable(client: TestClient, db_session: Session) -> None:
    _schedule_fixture(db_session)
    csrf = _csrf(_login(client, "adm"))
    client.post(
        "/work-schedule/entries",
        json={
            "entry_date": _iso(MONDAY),
            "time_from": "13:00",
            "time_to": "14:00",
            "title": "Увольнение 13:00–14:00",
        },
        headers={"X-CSRF-Token": csrf},
    )
    client.post(
        "/work-schedule/entries",
        json={"entry_date": _iso(TUESDAY), "title": "перевод"},
        headers={"X-CSRF-Token": csrf},
    )

    query = f"from={_iso(MONDAY)}&to={_iso(TUESDAY)}"
    api_items = client.get(f"/work-schedule?{query}").json()["items"]
    export = client.get(f"/work-schedule/export.xlsx?{query}")
    assert export.status_code == 200
    assert export.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert (
        'filename="work-schedule_2026-08-10_2026-08-11.xlsx"'
        in export.headers["content-disposition"]
    )

    workbook = load_workbook(io.BytesIO(export.content))
    sheet = _sheet(workbook)
    assert sheet.title == "График выхода"
    assert sheet.page_setup.orientation == "landscape"
    assert sheet.print_title_rows == "$4:$4"  # шапка повторяется на каждой странице
    assert sheet.freeze_panes == "A5"

    rows = list(sheet.iter_rows(values_only=True))
    day_headers = [row[0] for row in rows if isinstance(row[0], str) and "авг." in row[0]]
    assert day_headers == [
        "пн, 10 авг. 2026 — выходов: 3",
        "вт, 11 авг. 2026 — выходов: 1",
    ]
    assert rows[3][:3] == ("№", "ФИО", "Время")

    # Те же строки, что в API (номер, ФИО, время) и в том же порядке.
    exported = [(row[0], row[1], row[2]) for row in rows if isinstance(row[0], int)]
    expected = [
        (
            item["number"],
            item["display_name"],
            item["start_time"][:5] if item["start_time"] else "—",
        )
        for item in api_items
    ]
    assert exported == expected
    assert exported[0] == (1, "Антонов Антон", "08:00")

    # Экспорт попадает в аудит.
    assert (
        db_session.scalar(
            select(AuditEvent).where(AuditEvent.action == AuditAction.WORK_SCHEDULE_EXPORTED)
        )
        is not None
    )


def test_export_escapes_formula_injection(client: TestClient, db_session: Session) -> None:
    hr = make_user(db_session, username="hr1", role=UserRole.HR)
    csrf = _csrf(_login(client, "hr1"))
    candidate = make_candidate(db_session, owner=hr, stage=CandidateStage.REJECTED)
    _schedule(
        db_session,
        candidate,
        start_time="09:00",
        organization="=1+1",
        comment="=SUM(A1:A9)",
    )
    client.post(
        "/work-schedule/entries",
        json={"entry_date": _iso(MONDAY), "title": "@перевод", "comment": "+cmd"},
        headers={"X-CSRF-Token": csrf},
    )

    export = client.get(
        f"/work-schedule/export.xlsx?from={_iso(MONDAY)}&to={_iso(MONDAY)}&include_rejected=true"
    )
    assert export.status_code == 200
    sheet = _sheet(load_workbook(io.BytesIO(export.content)))
    cells = {(row[3], row[7]) for row in sheet.iter_rows(values_only=True) if row[1] is not None}
    values = [value for pair in cells for value in pair if isinstance(value, str)]
    assert "=1+1" in values and "=SUM(A1:A9)" in values
    for row in sheet.iter_rows(min_row=5):
        for cell in row:
            # Ни одно значение не стало формулой.
            assert cell.data_type != "f"
    # Заголовки дня и служебная строка остаются читаемым текстом.
    text = "\n".join(
        str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None
    )
    assert "@перевод" in text


def test_export_requires_authentication(client: TestClient, db_session: Session) -> None:
    make_user(db_session, username="hr1", role=UserRole.HR)
    assert client.get("/work-schedule/export.xlsx").status_code == 401


# --- Миграция ----------------------------------------------------------------


def test_migration_0016_is_the_head_and_adds_schedule_objects() -> None:
    """Ревизия 0016 — head цепочки и создаёт объекты графика.

    Само применение (upgrade/downgrade на PostgreSQL) проверяет integration
    job; здесь — структурная проверка, доступная и на SQLite-only прогоне.
    """

    from importlib.util import module_from_spec, spec_from_file_location
    from pathlib import Path

    from alembic.config import Config
    from alembic.script import ScriptDirectory

    migration_path = (
        Path(__file__).resolve().parents[1] / "alembic" / "versions" / "0016_work_schedule.py"
    )
    spec = spec_from_file_location("phase18_migration", migration_path)
    assert spec is not None and spec.loader is not None
    module = module_from_spec(spec)
    spec.loader.exec_module(module)
    assert module.revision == "0016"
    assert module.down_revision == "0015"

    config = Config("alembic.ini")
    script = ScriptDirectory.from_config(config)
    assert script.get_current_head() == "0016"

    from inspect import getsource

    source = getsource(module)
    for expected in (
        "schedule_entries",
        "start_date",
        "start_time",
        "start_organization",
        "start_department",
        "shift",
        "start_comment",
        "ix_candidates_start_date",
    ):
        assert expected in source


def test_schedule_rows_keep_dates_within_the_period(
    client: TestClient, db_session: Session
) -> None:
    """Границы периода включительные, дни идут по возрастанию."""

    hr = make_user(db_session, username="hr1", role=UserRole.HR)
    candidate = make_candidate(db_session, owner=hr, stage=CandidateStage.HIRED)
    _schedule(db_session, candidate, start_date=MONDAY)
    _login(client, "hr1")

    inside = client.get(f"/work-schedule?from={_iso(MONDAY)}&to={_iso(MONDAY)}").json()
    assert inside["total"] == 1
    assert inside["period_from"] == _iso(MONDAY)
    assert inside["period_to"] == _iso(MONDAY)

    after = client.get(f"/work-schedule?from={_iso(MONDAY + timedelta(days=1))}").json()
    assert after["total"] == 0

    open_ended = client.get(f"/work-schedule?to={_iso(MONDAY + timedelta(days=1))}").json()
    assert open_ended["total"] == 1
    assert open_ended["period_from"] is None
    assert UUID(inside["items"][0]["id"]) == candidate.id
