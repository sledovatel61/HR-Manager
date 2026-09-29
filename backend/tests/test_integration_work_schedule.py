"""Integration mirror of the work-schedule tests against a real PostgreSQL.

Run with::

    TEST_DATABASE_URL=postgresql+psycopg://user:pass@host:5432/db \
        pytest -m integration -v

Проверяется то, что специфично для PostgreSQL: native UUID/Date/Time колонки
новой миграции ``0016``, серверная выборка и выгрузка .xlsx на реальной БД,
а также аудит переноса даты (``audit_log.candidate_id`` + enum-значения).
"""

from __future__ import annotations

import io
from collections.abc import Iterator
from datetime import date
from typing import cast

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import AuditAction, AuditEvent, Candidate, CandidateStage, UserRole
from app.routers.auth import reset_login_limiter
from tests.conftest import FIXTURE_PASSWORD, make_candidate, make_user

pytestmark = pytest.mark.integration

MONDAY = date(2026, 8, 10)


@pytest.fixture(autouse=True)
def _clean_limiter() -> Iterator[None]:
    reset_login_limiter()
    yield
    reset_login_limiter()


def _login(client: TestClient, username: str) -> httpx.Response:
    return client.post("/auth/login", json={"username": username, "password": FIXTURE_PASSWORD})


def _schedule(db: Session, candidate: Candidate, **fields: object) -> Candidate:
    for name, value in fields.items():
        setattr(candidate, name, value)
    db.commit()
    db.refresh(candidate)
    return candidate


def test_start_date_round_trips_on_postgres(pg_client: TestClient, pg_db: Session) -> None:
    hr = make_user(pg_db, username="hr1", role=UserRole.HR)
    candidate = make_candidate(pg_db, owner=hr, stage=CandidateStage.OFFER)
    csrf = _login(pg_client, "hr1").json()["csrf_token"]

    saved = pg_client.patch(
        f"/candidates/{candidate.id}",
        json={
            "start_date": MONDAY.isoformat(),
            "start_time": "09:15",
            "start_organization": "ООО Авион",
            "start_department": "Производственный цех Сокол",
            "shift": "2 смена",
            "start_comment": "проходит медосмотр 13-го",
        },
        headers={"X-CSRF-Token": csrf},
    )
    assert saved.status_code == 200
    assert saved.json()["start_date"] == "2026-08-10"
    assert saved.json()["start_time"] == "09:15:00"

    listing = pg_client.get("/work-schedule?from=2026-08-01&to=2026-08-31")
    assert listing.status_code == 200
    body = listing.json()
    assert body["total"] == 1
    assert body["items"][0]["display_name"] == "Иванов Иван Иванович"
    assert body["items"][0]["organization"] == "ООО Авион"

    audit = pg_db.scalar(
        select(AuditEvent).where(AuditEvent.action == AuditAction.CANDIDATE_START_SCHEDULE_CHANGED)
    )
    assert audit is not None and audit.candidate_id == candidate.id
    assert "start_date: — -> 2026-08-10" in (audit.details or "")


def test_started_without_date_is_rejected_on_postgres(
    pg_client: TestClient, pg_db: Session
) -> None:
    hr = make_user(pg_db, username="hr1", role=UserRole.HR)
    candidate = make_candidate(pg_db, owner=hr, stage=CandidateStage.HIRED)
    csrf = _login(pg_client, "hr1").json()["csrf_token"]

    response = pg_client.patch(
        f"/candidates/{candidate.id}", json={"stage": "started"}, headers={"X-CSRF-Token": csrf}
    )
    assert response.status_code == 422
    assert "дату выхода" in response.json()["detail"].lower()


def test_schedule_scope_and_entries_on_postgres(pg_client: TestClient, pg_db: Session) -> None:
    hr1 = make_user(pg_db, username="hr1", role=UserRole.HR)
    hr2 = make_user(pg_db, username="hr2", role=UserRole.HR)
    mine = make_candidate(pg_db, owner=hr1, full_name="Свой Кандидат")
    _schedule(pg_db, mine, start_date=MONDAY, start_time=None)
    foreign = make_candidate(pg_db, owner=hr2, full_name="Чужой Кандидат")
    _schedule(pg_db, foreign, start_date=MONDAY)

    csrf = _login(pg_client, "hr1").json()["csrf_token"]
    created = pg_client.post(
        "/work-schedule/entries",
        json={
            "entry_date": MONDAY.isoformat(),
            "time_from": "13:00",
            "time_to": "14:00",
            "title": "Увольнение",
        },
        headers={"X-CSRF-Token": csrf},
    )
    assert created.status_code == 201

    listing = pg_client.get("/work-schedule?from=2026-08-01&to=2026-08-31").json()
    assert {item["kind"] for item in listing["items"]} == {"candidate", "entry"}
    # Строки без времени — в конце дня: служебная строка 13:00 идёт первой.
    assert [item["display_name"] for item in listing["items"]] == ["Увольнение", "Свой Кандидат"]

    export = pg_client.get("/work-schedule/export.xlsx?from=2026-08-01&to=2026-08-31")
    assert export.status_code == 200
    sheet = cast(Worksheet, load_workbook(io.BytesIO(export.content)).active)
    values = [row[1] for row in sheet.iter_rows(min_row=5, values_only=True) if row[1]]
    assert values == ["Увольнение", "Свой Кандидат"]

    deleted = pg_client.delete(
        f"/work-schedule/entries/{created.json()['id']}", headers={"X-CSRF-Token": csrf}
    )
    assert deleted.status_code == 204

    # Руководитель видит обе строки.
    _login(pg_client, "hr2")
    assert pg_client.get("/work-schedule?from=2026-08-01&to=2026-08-31").json()["total"] == 1
    make_user(pg_db, username="mgr", role=UserRole.MANAGER)
    _login(pg_client, "mgr")
    assert pg_client.get("/work-schedule?from=2026-08-01&to=2026-08-31").json()["total"] == 2
